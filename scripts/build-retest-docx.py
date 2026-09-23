from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


SOURCE = Path(r"C:\Users\0.0\Documents\xwechat_files\wxid_64tiyyhavjf012_8683\msg\file\2026-09\打金店管理系统完整测试文档(1).docx")
OUTPUT = Path(r"D:\xiangmu-wenjian\dajin-system\docs\打金店管理系统完整测试文档-复测版-2026-09-17.docx")

RESULTS = {
    "SALE-006": "[P] 复测通过（2026-09-17）\n旧料抵扣100元，实收500元；商品毛利按600-100=500元，未重复冲减。",
    "SALE-007": "[P] 复测通过（已修复）\n储值200元+现金300元组合支付成功，余额1000元降至800元；金额改为0后输入框仍保留，须主动关闭才移除。",
    "VIS-001": "[P] 复测通过\n会员成交支付后自动生成1条“顾客成交3天回访”，单号、金额、会员和购买时间完整。",
    "VIS-002": "[P] 复测通过（已修复）\n记录保存至 visit_task；移动端“已完成”显示拨号结果、回访正文和完成时间。",
    "PROC-002": "[P] 复测通过\n加工中转已完成不校验尾款；仅确认取货时拦截未收尾款，返回409705。",
    "PROC-004": "[P] 复测通过\n留店余料写入1条旧料库存记录并抵扣加工应收；未付款前销售额、加工收入和商品毛利不变。",
    "PROC-005": "[P] 复测通过（已修复）\n收银端待取货新增“确认取货”；尾款结清后可不打印直接更新为已取货。",
    "PROC-007": "[P] 复测通过（已修复）\n移动端只发送收银端待打印任务，不再打开手机打印预览。",
    "PROC-014": "[P] 复测通过（已修复）\n会员消费记录显示时间、订单号、商品、支付方式和金额。",
    "RPT-004": "[P] 复测通过\n今天、昨天、本周、本月、上月、自定义均可切换；自定义2026-09-01至2026-09-17返回正确。",
    "RPT-005": "[P] 复测通过（已修复）\n接口返回金价区间分布，移动端已展示；本月真实数据返回6个区间。",
    "SYNC-004": "[P] 复测通过（已修复）\n权限增删后同一令牌可获取最新权限；广播到达后移动端自动刷新入口，撤权页面跳转无权限页。",
}


def set_run_font(run, size=8, bold=False, color=None):
    run.font.name = "Microsoft YaHei"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def set_cell_text(cell, text):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05
    lines = text.split("\n")
    for index, line in enumerate(lines):
        if index:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        set_run_font(run, size=7.5, bold=index == 0, color=(20, 115, 70) if index == 0 else (55, 65, 75))
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margin(cell, top=70, start=70, bottom=70, end=70):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)


def move_after(anchor, element):
    anchor.addnext(element)
    return element


doc = Document(SOURCE)
doc.paragraphs[0].text = "打金店管理系统完整测试文档 复测版"
doc.paragraphs[0].style = doc.styles["Title"]
for run in doc.paragraphs[0].runs:
    set_run_font(run, size=22, bold=True, color=(0, 0, 0))

doc.paragraphs[1].text = "管理端 收银端 移动端 后端与基础设施  问题用例复测"
doc.paragraphs[1].alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in doc.paragraphs[1].runs:
    set_run_font(run, size=10, color=(75, 85, 95))

info = doc.tables[0]
info.cell(1, 1).text = "V1.1 复测版"
info.cell(1, 3).text = "2026-09-17"
info.cell(3, 3).text = "自动化回归 + 真实 API/数据库断言 + 生产构建"
for row in info.rows:
    for cell in row.cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                set_run_font(run, size=8, bold=False)

updated = []
for table in doc.tables:
    for row in table.rows:
        prevent_row_split(row)
        case_id = row.cells[0].text.strip() if row.cells else ""
        if case_id not in RESULTS:
            continue
        result_cell = row.cells[6]
        set_cell_text(result_cell, RESULTS[case_id])
        shade_cell(result_cell, "EAF7F0")
        set_cell_margin(result_cell)
        updated.append(case_id)

missing = sorted(set(RESULTS) - set(updated))
if missing:
    raise RuntimeError(f"Missing result rows: {missing}")

# Build a compact retest summary, then move it directly after the document information table.
heading = doc.add_paragraph(style="Heading 2")
heading.add_run("问题用例复测结论")
for run in heading.runs:
    set_run_font(run, size=13, bold=True, color=(35, 45, 55))

summary = doc.add_paragraph(style="Body Text")
summary.add_run(
    "2026-09-17 对“结果”栏已填写问题的 12 条用例重新执行测试；结果栏仅为 [ ] 的用例按用户说明视为已测，本轮未重复执行。"
    "12 条均通过，其中 7 条在复测过程中完成修复并加入回归测试。临时商品、会员、订单、加工单、流水和权限数据已全部清理。"
)
summary.paragraph_format.space_after = Pt(6)
for run in summary.runs:
    set_run_font(run, size=9, color=(45, 55, 65))

stats = doc.add_table(rows=2, cols=5)
stats.style = "Table Grid"
headers = ["复测用例", "通过", "失败", "修复项", "阻塞"]
values = ["12", "12", "0", "7", "0"]
for col, value in enumerate(headers):
    cell = stats.cell(0, col)
    cell.text = value
    shade_cell(cell, "34495E")
    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in paragraph.runs:
            set_run_font(run, size=8, bold=True, color=(255, 255, 255))
for col, value in enumerate(values):
    cell = stats.cell(1, col)
    cell.text = value
    shade_cell(cell, "F3F6F8" if col != 1 else "EAF7F0")
    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in paragraph.runs:
            set_run_font(run, size=10, bold=True, color=(20, 115, 70) if col == 1 else (35, 45, 55))
    set_cell_margin(cell, top=55, bottom=55)

evidence = doc.add_paragraph(style="Body Text")
evidence.add_run("复测证据：")
evidence.add_run(
    "后端 32/32、管理端 12/12、收银端 26/26、移动端 17/17 自动化测试通过；四端生产构建通过；"
    "真实 HTTP、数据库和状态流转测试通过；后端、MySQL、Redis、MinIO 健康。"
)
evidence.paragraph_format.space_before = Pt(5)
evidence.paragraph_format.space_after = Pt(8)
for index, run in enumerate(evidence.runs):
    set_run_font(run, size=8.5, bold=index == 0, color=(35, 45, 55))

anchor = info._tbl
for element in (heading._p, summary._p, stats._tbl, evidence._p):
    anchor = move_after(anchor, element)

# Add a short execution note at the end for traceability.
doc.add_heading("复测执行记录", level=1)
record = doc.add_paragraph(style="Body Text")
record.add_run("执行日期：").bold = True
record.add_run("2026-09-17\n")
record.add_run("复测范围：").bold = True
record.add_run("SALE-006、SALE-007、VIS-001、VIS-002、PROC-002、PROC-004、PROC-005、PROC-007、PROC-014、RPT-004、RPT-005、SYNC-004\n")
record.add_run("执行方式：").bold = True
record.add_run("自动化回归、真实接口请求、数据库前后值断言、加工状态流转、临时权限增删、生产构建。\n")
record.add_run("数据处理：").bold = True
record.add_run("真实业务复测使用隔离临时数据；执行结束后清理并核对残留数为 0。")
record.paragraph_format.space_after = Pt(6)
for run in record.runs:
    set_run_font(run, size=9, bold=run.bold, color=(45, 55, 65))

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
